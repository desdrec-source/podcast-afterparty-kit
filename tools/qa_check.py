#!/usr/bin/env python3
"""Episode QA checker. Usage: qa_check.py EPISODE_FOLDER [--lufs -16] [--tol 1] [--tp -1] [--forbid REGEX]
Run with the diarisation venv: /workspace/.diar/bin/python tools/qa_check.py FOLDER
Reads local .docx/.md/.txt/.html/.srt files. Exit code 0 if all PASS, 1 otherwise.
"""
import argparse, glob, os, re, subprocess, sys, zipfile
NOTICE = "This is an example draft for ideas and inspiration."
FOOTER = "Created for AI Video Motion members"
TEXT_EXT = (".docx", ".md", ".markdown", ".txt", ".html", ".htm")

def read_text(p):
    if p.endswith(".docx"):
        import docx
        d = docx.Document(p)
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for r in t.rows: parts += [c.text for c in r.cells]
        for s in d.sections:
            parts += [x.text for x in s.header.paragraphs + s.footer.paragraphs]
        return "\n".join(parts)
    t = open(p, encoding="utf-8", errors="ignore").read()
    if p.endswith((".html", ".htm")):
        import html; t = html.unescape(re.sub(r"<[^>]+>", "\n", t))
    return t

def docx_meta(p):
    try:
        with zipfile.ZipFile(p) as z: return " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.startswith("docProps/"))
    except Exception: return ""

def lines(t): return [l.strip().lstrip("#*-• ").strip() for l in t.splitlines() if l.strip()]

def to_s(s):
    v = 0
    for x in s.split(":"): v = v * 60 + int(x)
    return v

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("folder")
    ap.add_argument("--lufs", type=float, default=-16); ap.add_argument("--tol", type=float, default=1); ap.add_argument("--tp", type=float, default=-1)
    ap.add_argument("--forbid", default="", help="regex of names that must not appear anywhere (e.g. the owner's name)")
    a = ap.parse_args(); F = a.folder
    R = []
    def add(name, ok, detail): R.append((name, "PASS" if ok else "FAIL", detail))
    files = sorted(glob.glob(os.path.join(F, "**", "*"), recursive=True))
    docs = [f for f in files if f.lower().endswith(TEXT_EXT) and not os.path.basename(f).startswith("~$")]
    T = {f: read_text(f) for f in docs}
    pick = lambda pat: [f for f in docs if re.search(pat, os.path.basename(f), re.I)]
    sn = sorted(pick(r"show.?notes"), key=lambda f: not f.endswith(".docx"))

    # show notes title + one sentence
    if sn:
        L = lines(T[sn[0]]); title = L[0] if L else ""
        add("Show-notes title < 60 chars", 0 < len(title) < 60, f"{len(title)}: {title[:60]}")
        one = next((l for l in L if re.match(r"(in one sentence|one.sentence answer|tl;?dr)\s*:", l, re.I)), None)
        if one: one = one.split(":", 1)[1].strip()
        add("One-sentence answer <= 150 chars", bool(one) and len(one) <= 150, f"{len(one)} chars" if one else "not found ('In one sentence:' line)")
    else:
        add("Show-notes title < 60 chars", False, "no show-notes file"); add("One-sentence answer <= 150 chars", False, "no show-notes file")

    # chapters: from show notes (fallback any doc), section after 'Chapters' heading
    src = sn + pick(r"description") if sn else docs
    ch, chsrc = [], ""
    for f in src:
        L = lines(T[f]); got = []; on = False
        for l in L:
            if re.fullmatch(r"chapters?:?", l, re.I): on = True; continue
            if on:
                m = re.match(r"\(?(\d{1,2}:\d{2}(?::\d{2})?)\)?\s*[-–—]?\s*(.+)", l)
                if m: got.append((to_s(m[1]), m[2]))
                elif got: break
        if got: ch, chsrc = got, os.path.basename(f); break
    starts = [c[0] for c in ch]
    probs = []
    if not ch: probs.append("no chapters found")
    else:
        if starts[0] != 0: probs.append("first != 00:00")
        if len(ch) < 3: probs.append(f"only {len(ch)}")
        for x, y in zip(starts, starts[1:]):
            if y <= x: probs.append(f"not ascending at {y}s")
            elif y - x < 10: probs.append(f"chapter at {x}s < 10s")
    add("Chapters (00:00, >=3, ascending, >=10s)", not probs, f"{len(ch)} in {chsrc}" + ("; " + "; ".join(probs) if probs else ""))

    # notice + footer in written docs (all text docs except transcripts/captions/style guide); both required
    written = [f for f in docs if not re.search(r"transcript|caption|\.srt$|style.?guide", os.path.basename(f), re.I)]
    for label, needle in (("Example-draft notice in written docs", NOTICE), ("AI Video Motion footer in written docs", FOOTER)):
        miss = [os.path.basename(f) for f in written if needle not in " ".join(T[f].split())]
        add(label, bool(written) and not miss, (f"{len(written)-len(miss)}/{len(written)} have it" if written else "no written docs") + (f"; missing: {', '.join(miss)}" if miss else ""))

    # optional: forbidden names anywhere (text, file names, docx metadata, srt, mp3 tags)
    if a.forbid:
        pat = re.compile(a.forbid, re.I)
        hits = []
        for f in files:
            if os.path.isdir(f): continue
            bag = os.path.relpath(f, F)
            if f in T: bag += "\n" + T[f]
            if f.endswith(".srt"): bag += "\n" + open(f, encoding="utf-8", errors="ignore").read()
            if f.endswith(".docx"): bag += "\n" + docx_meta(f)
            if f.lower().endswith(".mp3"):
                bag += subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format_tags", "-of", "default=nw=1", f], capture_output=True, text=True).stdout
            for m in pat.finditer(bag): hits.append(f"{os.path.basename(f)}:'{bag[max(0,m.start()-15):m.end()+15]!s}'".replace("\n", " "))
        add("No forbidden names anywhere", not hits, "none" if not hits else "; ".join(hits[:5]))

    # MP3 loudness (the normalised file; not 'original')
    mp3s = [f for f in files if f.lower().endswith(".mp3") and "original" not in os.path.basename(f).lower()]
    if mp3s:
        out = subprocess.run(["ffmpeg", "-hide_banner", "-nostats", "-i", mp3s[0], "-af", "ebur128=peak=true:framelog=quiet", "-f", "null", "-"], capture_output=True, text=True).stderr.split("Summary:")[-1]
        g = lambda k: float(re.search(k + r":\s+(-?[\d.]+)", out).group(1))
        I, TP, LRA = g("I"), g("Peak"), g("LRA")
        add("MP3 loudness in spec", abs(I - a.lufs) <= a.tol and TP <= a.tp, f"{os.path.basename(mp3s[0])[:40]}: I={I} TP={TP} LRA={LRA}")
    else: add("MP3 loudness in spec", False, "no non-original MP3")

    # images
    from PIL import Image
    imgs = [f for f in files if f.lower().endswith((".jpg", ".jpeg", ".png"))]
    def img(pat, size, strict):
        c = [f for f in imgs if re.search(pat, os.path.basename(f), re.I)]
        if not c: return False, "not found"
        im = Image.open(c[0]); ok = im.size == size and (not strict or (im.format == "JPEG" and im.mode == "RGB"))
        return ok, f"{os.path.basename(c[0])[:40]}: {im.format} {im.mode} {im.size[0]}x{im.size[1]}"
    add("Cover 3000x3000 RGB JPEG", *img(r"cover", (3000, 3000), True))
    add("Thumbnail 1280x720", *img(r"thumb", (1280, 720), False))

    # SRT
    srts = [f for f in files if f.lower().endswith(".srt")]
    if srts:
        t = open(srts[0], encoding="utf-8-sig", errors="ignore").read().replace("\r", "")
        blocks = [b for b in re.split(r"\n\s*\n", t.strip()) if b.strip()]
        errs, prev, n = [], -1, 0
        for i, b in enumerate(blocks, 1):
            L = b.split("\n")
            m = re.fullmatch(r"(\d\d):(\d\d):(\d\d),(\d{3}) --> (\d\d):(\d\d):(\d\d),(\d{3})", L[1].strip()) if len(L) > 1 else None
            if not L[0].strip().isdigit() or not m: errs.append(f"block {i}: bad index/timing"); continue
            s = int(m[1])*3600+int(m[2])*60+int(m[3])+int(m[4])/1000; e = int(m[5])*3600+int(m[6])*60+int(m[7])+int(m[8])/1000
            if e <= s: errs.append(f"#{L[0]} end<=start")
            if s < prev - 0.001: errs.append(f"#{L[0]} overlaps/out of order")
            prev = e; txt = L[2:]
            if not txt: errs.append(f"#{L[0]} empty")
            if len(txt) > 2: errs.append(f"#{L[0]} {len(txt)} lines")
            for x in txt:
                if len(x) > 42: errs.append(f"#{L[0]} line {len(x)} chars"); break
            n += 1
        add("SRT valid, <=2 lines x 42 chars", not errs, f"{n} cues" + (f"; {len(errs)} issues: " + "; ".join(errs[:4]) if errs else ""))
    else: add("SRT valid, <=2 lines x 42 chars", False, "no .srt")

    w = max(len(r[0]) for r in R)
    print(f"QA: {F}\n" + "-" * (w + 60))
    for n, s, d in R: print(f"{n:<{w}}  {s:<4}  {d}")
    fails = sum(r[1] == "FAIL" for r in R)
    print("-" * (w + 60) + f"\n{len(R)-fails}/{len(R)} PASS" + ("" if not fails else f", {fails} FAIL"))
    sys.exit(1 if fails else 0)

if __name__ == "__main__": main()
